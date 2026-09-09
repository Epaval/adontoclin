import io
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db import models
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.views import View
from django.views.generic import TemplateView

from apps.billing.models import DetalleFactura, Factura
from apps.clinical.models import ServicioDental

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    openpyxl = None


class SoloSuperUser(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_superuser


class LicenciaView(SoloSuperUser, TemplateView):
    template_name = "core/licencia_panel.html"

    def get_context_data(self, **kwargs):
        from apps.core import licencias as L

        def _call(nombre, default):
            fn = getattr(L, nombre, None)
            if fn is None:
                return default
            try:
                v = fn()
                return v if v not in (None, "") else default
            except Exception:
                return default

        ctx = super().get_context_data(**kwargs)
        ctx.update({
            "estado": _call("estado_licencia", "SIN CLAVE · modo prueba"),
            "tipo": _call("tipo_licencia", "Prueba"),
            "dias": _call("dias_restantes", "—"),
            "vencimiento": _call("fecha_vencimiento", "—"),
            "huella": _call("huella_maquina", "—"),
            "clave": _call("clave_activada", "— (sin clave: modo prueba)"),
            "modo_dev": _call("modo_dev", False),
        })
        return ctx



class EstadisticasView(SoloSuperUser, TemplateView):
    template_name = "reports/estadisticas.html"

    def get_context_data(self, **kwargs):
        import json
        from datetime import date, timedelta

        from django.db.models import Count, Sum

        from apps.billing.models import DetalleFactura, Factura
        from apps.clinical.models import CitaDental, HistorialDiente
        from apps.patients.models import Paciente

        ctx = super().get_context_data(**kwargs)
        hoy = date.today()

        top = list(DetalleFactura.objects.values("servicio__nombre")
                   .annotate(total=Count("id")).order_by("-total")[:10])
        ctx["top_labels"] = json.dumps([t["servicio__nombre"] for t in top])
        ctx["top_data"] = json.dumps([t["total"] for t in top])

        meses, ingresos = [], []
        for i in range(11, -1, -1):
            f = hoy.replace(day=1) - timedelta(days=31 * i)
            mes = f.replace(day=1)
            sig = (mes.replace(day=28) + timedelta(days=4)).replace(day=1)
            t = Factura.objects.filter(estado="emitida", fecha_creacion__gte=mes, fecha_creacion__lt=sig).aggregate(t=Sum("total"))["t"] or 0
            meses.append(mes.strftime("%m/%y"))
            ingresos.append(float(t))
        ctx["mes_labels"] = json.dumps(meses)
        ctx["mes_data"] = json.dumps(ingresos)

        sexo = list(Paciente.objects.filter(activo=True).values("sexo").annotate(total=Count("id")))
        ctx["sexo_labels"] = json.dumps([s["sexo"] for s in sexo])
        ctx["sexo_data"] = json.dumps([s["total"] for s in sexo])

        rangos = {"0-12": 0, "13-17": 0, "18-35": 0, "36-55": 0, "56+": 0}
        for p in Paciente.objects.filter(activo=True):
            e = p.edad
            k = "0-12" if e <= 12 else "13-17" if e <= 17 else "18-35" if e <= 35 else "36-55" if e <= 55 else "56+"
            rangos[k] += 1
        ctx["edad_labels"] = json.dumps(list(rangos.keys()))
        ctx["edad_data"] = json.dumps(list(rangos.values()))

        trat = list(HistorialDiente.objects.values("tipo").annotate(total=Count("id")).order_by("-total"))
        ctx["trat_labels"] = json.dumps([t["tipo"] for t in trat])
        ctx["trat_data"] = json.dumps([t["total"] for t in trat])

        estados = list(CitaDental.objects.values("estado").annotate(total=Count("id")))
        ctx["estado_labels"] = json.dumps([e["estado"] for e in estados])
        ctx["estado_data"] = json.dumps([e["total"] for e in estados])
        return ctx



class AjusteMasivoPreciosView(SoloSuperUser, View):
    """Ajuste masivo de precios de servicios en $ por porcentaje."""

    def get(self, request):
        return render(request, "clinical/ajuste_masivo.html", {
            "servicios": ServicioDental.objects.filter(activo=True),
        })

    def post(self, request):
        try:
            pct = Decimal(request.POST.get("porcentaje", "0"))
        except Exception:
            messages.error(request, "Porcentaje inválido.")
            return redirect("ajuste_masivo")
        if pct == 0:
            messages.error(request, "Indique un porcentaje distinto de 0.")
            return redirect("ajuste_masivo")
        factor = (Decimal("100") + pct) / Decimal("100")
        n = 0
        for s in ServicioDental.objects.filter(activo=True):
            s.precio_usd = (s.precio_usd * factor).quantize(Decimal("0.01"))
            s.save(update_fields=["precio_usd"])
            n += 1
        messages.success(request, f"{n} servicios ajustados en {pct}%.")
        return redirect("ajuste_masivo")


class FacturaExportarExcelView(SoloSuperUser, TemplateView):
    template_name = "reports/exportar_excel.html"

    def get(self, request, *args, **kwargs):
        if request.GET.get("descargar") == "1":
            return self.post(request, *args, **kwargs)
        return super().get(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        import datetime as dt
        import os
        import os as _os

        from django.conf import settings

        qs = Factura.objects.select_related("cita__expediente__paciente").order_by("-fecha_creacion")
        desde = request.POST.get("desde") or request.GET.get("desde")
        hasta = request.POST.get("hasta") or request.GET.get("hasta")
        if desde:
            qs = qs.filter(fecha_creacion__date__gte=desde)
        if hasta:
            qs = qs.filter(fecha_creacion__date__lte=hasta)

        def _v(x):
            return x() if callable(x) else x

        filas = []
        for f in qs:
            pac = f.cita.expediente.paciente
            filas.append([
                f.numero, f.numero_control, f.fecha_creacion.strftime("%d/%m/%Y %H:%M"),
                _v(pac.full_name),
                _v(pac.ci_efectivo) or "-",
                f.get_estado_display(), f.get_metodo_pago_display() if hasattr(f, "get_metodo_pago_display") else "",
                float(f.subtotal), float(f.descuento), float(f.total), float(f.tasa),
                round(float(f.total) * float(f.tasa), 2),
            ])
        encabezado = ["Numero", "Control", "Fecha", "Paciente", "CI", "Estado", "Metodo",
                      "Subtotal $", "Descuento $", "Total $", "Tasa Bs/$", "Total Bs"]

        log_dir = os.path.join(getattr(settings, "DATA_DIR", settings.BASE_DIR), "logs")
        os.makedirs(log_dir, exist_ok=True)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "Facturacion"
            ws.append(encabezado)
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="0369A1")
            for fila in filas:
                ws.append(fila)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16
            import io as _io
            from django.conf import settings as _st
            if getattr(_st, "MODO_ESCRITORIO", False) or _os.environ.get("LABCLIN_MODO") == "escritorio":
                import os as _oso
                destino = _oso.path.join(_oso.path.expanduser("~"), "Documents", "OdontoClin")
                _oso.makedirs(destino, exist_ok=True)
                nombre = f"facturacion_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
                ruta = _oso.path.join(destino, nombre)
                wb.save(ruta)
                from django.shortcuts import render as _render
                return _render(request, "reports/exportado_local.html", {"ruta": ruta})
            buf = _io.BytesIO()
            wb.save(buf)
            response = HttpResponse(
                buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = 'attachment; filename="facturacion_odontoclin.xlsx"'
            return response
        except Exception as exc:
            # excel_fallback_csv: nunca dejar sin descarga
            with open(os.path.join(log_dir, "excel.log"), "a", encoding="utf-8") as lf:
                lf.write(f"{dt.datetime.now()}: {exc}\n")
            import csv as _csv
            import io as _io2
            buf2 = _io2.StringIO()
            w = _csv.writer(buf2, delimiter=";")
            w.writerow(encabezado)
            w.writerows(filas)
            response = HttpResponse(buf2.getvalue().encode("utf-8-sig"), content_type="text/csv")
            response["Content-Disposition"] = 'attachment; filename="facturacion_odontoclin.csv"'
            return response




class AbrirCarpetaExportView(LoginRequiredMixin, View):
    def get(self, request):
        import os
        import subprocess
        import sys
        destino = os.path.join(os.path.expanduser("~"), "Documents", "OdontoClin")
        os.makedirs(destino, exist_ok=True)
        try:
            if sys.platform == "win32":
                os.startfile(destino)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", destino])
            else:
                subprocess.Popen(["xdg-open", destino])
        except Exception:
            pass
        return redirect("exportar_excel")


class AdminAccesoRemotoView(LoginRequiredMixin, TemplateView):
    """Panel de acceso remoto: URL del túnel + estado + QR."""
    template_name = "admin/acceso_remoto.html"

    def get_context_data(self, **kwargs):
        import subprocess
        import requests
        import qrcode
        import io
        import base64
        
        ctx = super().get_context_data(**kwargs)
        url_tunel = "https://demo.facdin.com"
        ctx["url_tunel"] = url_tunel
        
        # Verificar estado del túnel
        estado = "desconocido"
        try:
            r = requests.get(url_tunel + "/accounts/login/", timeout=5)
            if r.status_code in [200, 302]:
                estado = "activo"
            else:
                estado = "caido"
        except Exception:
            estado = "caido"
        ctx["estado_tunel"] = estado
        
        # Generar QR
        qr = qrcode.QRCode(version=1, box_size=10, border=4)
        qr.add_data(url_tunel)
        qr.make(fit=True)
        img = qr.make_image(fill_color="#1e293b", back_color="white")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        qr_b64 = base64.b64encode(buf.getvalue()).decode()
        ctx["qr_base64"] = qr_b64
        
        return ctx
